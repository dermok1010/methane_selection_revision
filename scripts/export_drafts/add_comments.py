#!/usr/bin/env python
"""Re-attach a live Claude Doc's comment threads to its Word export as Word comments.

The Docs export drops comments, and build.py builds the .docx from markdown, so the
threads (Nóirín's review comments, Claude's replies, change notes) are added back here.
Each thread becomes one Word comment on the same words it was anchored to in the doc;
replies follow as extra paragraphs prefixed with their author. Resolved threads are kept and marked "[Resolved]". Threads whose anchor
text was later deleted are placed where that text stood (found from the character ids
around it) and marked as moved.

Inputs come from the docs tools: the tab's prose view (`read` of the body node, which
marks each thread's start with <comment id=…/>) and its `query` rows.

Usage (docx-tools env):
  python add_comments.py IN.docx VIEW.xml COMMENTS.json OUT.docx
"""
import copy, json, re, sys
import xml.etree.ElementTree as ET

import docx


def author_text(body):
    for pre, who in (("Nóirín McHugh:", "Nóirín McHugh"), ("Claude:", "Claude")):
        if body.startswith(pre):
            return who, body[len(pre):].strip()
    if body.startswith("Change note:"):
        return "Claude", body
    if body.startswith(("How the ", "Why a ")):  # explanatory notes Claude wrote for Dermot
        return "Claude", body
    return "Dermot Kelly", body.replace("@Claude⁠", "@Claude").strip()


def para_text(para):
    """Plain text of a view paragraph, the offset at which each thread starts, and the
    character-id start of each text run (to place threads whose anchor text was deleted)."""
    text, marks, runs = "", [], []
    for ch in para:
        if ch.tag == "comment":
            marks.append((ch.get("id"), len(text)))
        elif ch.tag == "text":
            t = "".join(ch.itertext())
            runs.append((ch.get("fullid"), len(text), len(t)))
            text += t
        elif ch.tag == "date":
            text += ch.get("value") or ""
        elif ch.tag == "mention":
            text += ch.get("name") or ""
    return text, marks, runs


def expand_ids(root):
    """Short ids ('.90') take the session of the last full id printed before them."""
    session = None
    for el in root.iter():
        i = el.get("id")
        if not i or "-" in i:  # comment ids are uuids
            continue
        if not i.startswith("."):
            session = i.split(".")[0]
        el.set("fullid", f"{session}.{i.split('.')[-1]}")


def near(text, pos, width=40):
    """The few words just before pos (or after it, at a paragraph start)."""
    before = text[max(0, pos - width):pos]
    if len(before.strip()) >= 6:
        return before[before.find(" ") + 1:] if " " in before[:-1] and pos > width else before
    return text[pos:pos + width]


def anchor_len(anchor):
    a, b = anchor.get("range", [None, None])
    if a and b and a.split(".")[0] == b.split(".")[0]:
        return int(b.split(".")[1]) - int(a.split(".")[1]) + 1
    return 30


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def isolate(par, a, b):
    """Split runs so chars [a, b) of the paragraph are whole runs; return them."""
    pos, out = 0, []
    for run in list(par.runs):
        t = run.text
        s, e = pos, pos + len(t)
        pos = e
        if e <= a or s >= b or not t:
            continue
        cuts = sorted({max(a, s) - s, min(b, e) - s} - {0, len(t)})
        pieces = [t[i:j] for i, j in zip([0] + cuts, cuts + [len(t)])]
        cur, ps = run, s
        for k, piece in enumerate(pieces):
            if k:
                new = copy.deepcopy(cur._r)
                cur._r.addnext(new)
                cur = docx.text.run.Run(new, par)
            cur.text = piece
            if ps >= a and ps + len(piece) <= b:
                out.append(cur)
            ps += len(piece)
    return out


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def find(doc, snippet, context):
    """Locate snippet in the document; fall back to shorter prefixes, then the paragraph."""
    paras = list(all_paragraphs(doc))
    for probe in (snippet, snippet[:25], snippet[:12]):
        probe = probe.strip()
        if len(probe) < 6:
            continue
        for par in paras:
            i = par.text.find(probe)
            if i >= 0:
                return par, i, i + len(probe)
    head = norm(context)[:40]
    for par in paras:
        if head and norm(par.text).startswith(head):
            return par, 0, min(len(par.text), 30)
    return None


def main(in_docx, view_xml, comments_json, out_docx):
    doc = docx.Document(in_docx)
    root = ET.fromstring(open(view_xml).read())
    raw = json.load(open(comments_json))
    resolved = {r["payload"]["value"]["parent"]["id"] for r in raw
                if r.get("verb") == "create" and r["payload"]["value"].get("kind") == "resolve"}
    rows = [r for r in raw if r.get("verb") == "create" and "body" in r["payload"]["value"]]
    by_id = {r["id"]: dict(r["payload"]["value"]) for r in rows}
    for r in raw:  # later edits to a comment's text
        if r.get("verb") == "update" and r["id"] in by_id:
            for op in r["payload"].get("patch", []):
                if op.get("op") == "set" and op.get("path") == ["body"]:
                    by_id[r["id"]]["body"] = op["value"]

    def thread(rid):
        out, queue = [], [rid]
        while queue:
            p = queue.pop(0)
            kids = [r["id"] for r in rows if r["payload"]["value"]["parent"].get("object") == "utterance"
                    and r["payload"]["value"]["parent"].get("id") == p]
            out += kids
            queue += kids
        return out

    expand_ids(root)
    placed, detached = [], []
    anchored, run_index = set(), []
    for para in root.iter("paragraph"):
        text, marks, runs = para_text(para)
        run_index += [(rid, text, off, n) for rid, off, n in runs if rid]
        for cid, off in marks:
            if cid not in by_id:
                continue
            anchored.add(cid)
            n = anchor_len(by_id[cid]["parent"].get("anchor", {}))
            snippet = text[off:off + n]
            if len(snippet.strip()) < 6:  # anchor collapsed to a point: use the words before it
                snippet = near(text, off)
            hit = find(doc, snippet, text)
            (placed.append((cid, hit, None)) if hit else detached.append(cid))
    roots = [r["id"] for r in rows if r["payload"]["value"]["parent"].get("object") == "node"]
    moved = "[Moved: the text this was anchored to was later deleted or rewritten; placed where it stood.]"
    for cid in [c for c in roots if c not in anchored] + detached:
        start = by_id[cid]["parent"].get("anchor", {}).get("range", [""])[0]
        sess, _, clock = start.rpartition(".")
        best = None
        for rid, text, off, n in run_index:
            rs, _, rc = rid.rpartition(".")
            if rs == sess and int(rc) <= int(clock or 0) and (best is None or int(rc) > int(best[0].rpartition(".")[2])):
                best = (rid, text, off, n)
        hit = None
        if best:
            rid, text, off, n = best
            pos = off + min(int(clock) - int(rid.rpartition(".")[2]), n)
            hit = find(doc, near(text, pos), text)
        placed.append((cid, hit, moved))
    detached = [cid for cid, hit, _ in placed if hit is None]

    def add(runs, cid, prefix=None):
        who, body = author_text(by_id[cid]["body"])
        lines = body.split("\n")
        tags = (["[Resolved]"] if cid in resolved else []) + ([prefix] if prefix else [])
        first = " ".join(tags + [lines[0]])
        c = doc.add_comment(runs, text=first, author=who, initials="".join(w[0] for w in who.split()))
        for ln in lines[1:]:
            c.add_paragraph(ln)
        for rid in thread(cid):
            w, t = author_text(by_id[rid]["body"])
            for k, ln in enumerate(t.split("\n")):
                c.add_paragraph((f"{w}: " if k == 0 else "") + ln)

    title = next(p for p in doc.paragraphs if p.text.strip())
    for cid, hit, prefix in placed:
        if hit is None:
            add(isolate(title, 0, len(title.text)), cid, prefix="[Detached: its anchor text could not be found.]")
        else:
            par, a, b = hit
            add(isolate(par, a, b), cid, prefix=prefix)
    doc.save(out_docx)
    n_moved = sum(1 for _, h, p in placed if p and h)
    print(f"threads {len(placed)}: in place {len(placed) - n_moved - len(detached)}, "
          f"moved to former position {n_moved}, on title {len(detached)}")
    return placed


if __name__ == "__main__":
    main(*sys.argv[1:5])
