#!/usr/bin/env python
"""Export the manuscript and response letter (live Claude Doc markdown exports) to
.docx in the submitted manuscript's own Word styles, with line numbers.

1. Manuscript: built on the submitted .docx (its styles.xml and page setup, continuous
   line numbering), mirroring draft 5's per-element formatting. Passages listed in
   spans.HIGHLIGHT are highlighted yellow.
2. The manuscript is rendered to PDF with LibreOffice and the line numbers of every
   span are read from it (pdftotext).
3. Response letter: same styles; each response gets the revised-manuscript line
   numbers of its spans (spans.LETTER_REFS).

Line numbers come from the LibreOffice rendering (Times New Roman -> metric-compatible
Liberation Serif), so the PDFs written alongside are the reference; Word's own
layout can differ by a line or so.

Usage (docx-tools env):
  python build.py MANUSCRIPT.md LETTER.md OUT_MANUSCRIPT.docx OUT_LETTER.docx
  python build.py --endnote endnote_cites.json MANUSCRIPT.md OUT_MANUSCRIPT.docx
"""
import copy, re, subprocess, sys, tempfile, unicodedata
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Emu

sys.path.insert(0, str(Path(__file__).parent))
import spans

REPO = Path(__file__).resolve().parents[2]
SUBMITTED = REPO / "manuscript/submitted/Methane_Genetic_Selection_Response_DK_Fnl.docx"
FIGURE1 = REPO / "analysis/revision/selection_index/results/tri_frontier_co2.png"
TOP_SECTIONS = {"Abstract", "Background", "Methods", "Results", "Discussion", "Conclusions", "Supplementary Material"}
GRID_TABLES = {"Table 4."}  # draft 5 / submitted use Table Grid here, Plain Table 2 elsewhere

# Plain-text renderings of the doc's LaTeX blocks, as in draft 5 (italic, 11.5 pt, centred).
# Each is a list of (text, vert) with vert in {None, "sub", "sup"}.
EQUATIONS = {
    r"het = 1 - \sum_{i=1}^{n} sire_i \times dam_i":
        [("het = 1 − Σ", None), ("i=1", "sub"), ("n", "sup"), (" sire", None), ("i", "sub"), (" × dam", None), ("i", "sub")],
    r"rec = 1 - \sum_{i=1}^{n} \frac{sire_i^2 + dam_i^2}{2}":
        [("rec = 1 − Σ", None), ("i=1", "sub"), ("n", "sup"), (" (sire", None), ("i", "sub"), ("2", "sup"),
         (" + dam", None), ("i", "sub"), ("2", "sup"), (") / 2", None)],
    r"y = \mu + sex + BR + CL + CV + LY + SU + TX + UN + het + rec + age + BT_g + RT_g + BT_e + RT_e + DP + CG + a + pe + e":
        [("y = μ + sex + BR + CL + CV + LY + SU + TX + UN + het + rec + age + BT", None), ("g", "sub"), (" + RT", None),
         ("g", "sub"), (" + BT", None), ("e", "sub"), (" + RT", None), ("e", "sub"), (" + DP + CG + a + pe + e", None)],
    r"h^2 = \frac{\sigma^2_a}{\sigma^2_a + \sigma^2_{pe} + \sigma^2_e}":
        [("h² = σ²ₐ / (σ²ₐ + σ²ₚₑ + σ²ₑ)", None)],
    r"t = \frac{\sigma^2_a + \sigma^2_{pe}}{\sigma^2_a + \sigma^2_{pe} + \sigma^2_e}":
        [("t = (σ²ₐ + σ²ₚₑ) / (σ²ₐ + σ²ₚₑ + σ²ₑ)", None)],
}


# ---------------------------------------------------------------- markdown helpers
def unescape(s):
    s = s.replace("&#32;", "").replace("&#91;", "[")
    return re.sub(r"\\([\[\]\*_~#()+\-.!`>|])", r"\1", s)


def inline(s):
    """Markdown inline -> list of segments dict(text, b, i, sup)."""
    s = s.replace("\\*", "\x00")
    segs, b, it = [], False, False
    for tok in re.split(r"(\*\*\*|\*\*|\*)", s):
        if tok == "***":
            b, it = not b, not it
        elif tok == "**":
            b = not b
        elif tok == "*":
            it = not it
        elif tok:
            t = unescape(tok.replace("\x00", "*"))
            # live weight^0.75 -> superscript
            parts = re.split(r"\^(0\.75)", t)
            for k, p in enumerate(parts):
                if p:
                    segs.append(dict(text=p, b=b, i=it, vert="sup" if k % 2 else None))
    return segs


def plain(segs):
    return "".join(s["text"] for s in segs)


def apply_highlights(segs, ids):
    """Split segments so that each highlight span carries hl=True. ids: span ids to try."""
    text = plain(segs)
    marks = [False] * len(text)
    found = []
    for sid in ids:
        start, end = spans.HIGHLIGHT[sid]
        a = text.find(start)
        if a < 0:
            continue
        if sid == "m_eq":  # last "+ e" of the model equation
            a = text.rfind(start)
        z = a + len(start) if end is None else text.find(end, a)
        if end is not None:
            if z < 0:
                continue
            z += len(end)
        for k in range(a, z):
            marks[k] = True
        found.append(sid)
    out, pos = [], 0
    for s in segs:
        t = s["text"]
        k = 0
        while k < len(t):
            m = marks[pos + k]
            j = k
            while j < len(t) and marks[pos + j] == m:
                j += 1
            out.append(dict(s, text=t[k:j], hl=m))
            k = j
        pos += len(t)
    return out, found


# ---------------------------------------------------------------- docx helpers
def fresh_document():
    d = docx.Document(str(SUBMITTED))
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    for el in list(body):
        if el is not sect:
            body.remove(el)
    # one portrait section with continuous line numbering, as in draft 5
    for tag in ("w:lnNumType",):
        for old in sect.findall(qn(tag)):
            sect.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    pgmar = sect.find(qn("w:pgMar"))
    pgmar.addnext(ln)
    pgsz = sect.find(qn("w:pgSz"))
    pgsz.set(qn("w:w"), "11906"); pgsz.set(qn("w:h"), "16838")
    if pgsz.get(qn("w:orient")):
        del pgsz.attrib[qn("w:orient")]
    return d


# EndNote mode (--endnote CITES.json): numeric citations become EndNote temporary
# citations ({Author, Year #RecNum}) and the typed reference list is omitted, so that
# "Update Citations and Bibliography" in Word relinks everything to the EndNote library.
CITES = None
CITE_RE = re.compile(r"\[(\d+(?:\s*[–-]\s*\d+)?(?:\s*,\s*\d+(?:\s*[–-]\s*\d+)?)*)\]")


def endnote_citations(text):
    def repl(m):
        nums = []
        for part in m.group(1).split(","):
            a, _, b = part.strip().replace("–", "-").partition("-")
            nums += list(range(int(a), int(b) + 1)) if b else [int(a)]
        missing = [n for n in nums if str(n) not in CITES]
        if missing:
            raise KeyError(f"no EndNote citation for reference(s) {missing} in {m.group(0)}")
        return "{" + "; ".join(CITES[str(n)] for n in nums) + "}"
    return CITE_RE.sub(repl, text)


def add_runs(p, segs, size=None, base_italic=False, highlight=False):
    for s in segs:
        if not s["text"]:
            continue
        r = p.add_run(endnote_citations(s["text"]) if CITES else s["text"])
        r.bold = bool(s.get("b"))
        r.italic = bool(s.get("i")) or base_italic
        if size:
            r.font.size = Pt(size)
        if s.get("vert") == "sup":
            r.font.superscript = True
        elif s.get("vert") == "sub":
            r.font.subscript = True
        if s.get("hl") or highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def para(d, segs, style=None, jc=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    if jc is not None:
        p.alignment = jc
    return add_runs(p, segs, **kw)


def heading_top(d, text, style=None):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    r = p.add_run(text)
    r.bold = True; r.italic = False; r.font.size = Pt(18)
    return p


def table(d, rows, style, hl_row_key=None):
    ncol = len(rows[0])
    t = d.add_table(rows=len(rows), cols=ncol)
    t.style = d.styles[style]
    for ri, row in enumerate(rows):
        hl = hl_row_key is not None and ri > 0 and hl_row_key in row[:2]
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if ci < 2 else WD_ALIGN_PARAGRAPH.CENTER
            segs = [dict(s, b=s["b"]) for s in inline(cell)]
            add_runs(p, segs, size=9.5, highlight=hl)
    return t


# ---------------------------------------------------------------- manuscript
def build_manuscript(md_path, out_path):
    lines = Path(md_path).read_text().splitlines()
    d = fresh_document()
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("```latex"):
            j = i + 1
            eq = []
            while not lines[j].startswith("```"):
                eq.append(lines[j]); j += 1
            blocks.append(("eq", "\n".join(eq).strip())); i = j + 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows)); continue
        m = re.match(r"(#+) (.*)", ln)
        if m:
            blocks.append(("h%d" % len(m.group(1)), unescape(m.group(2)))); i += 1; continue
        blocks.append(("p", ln)); i += 1

    eq_ids = {"m_eq", "m_h2", "m_t"}  # matched only inside equation blocks
    hl_ids = [k for k in spans.HIGHLIGHT if k not in eq_ids]
    found_all = set()
    found_count = {}
    section, last_caption, after_table, in_front = None, None, False, True
    title_done = False
    for kind, val in blocks:
        if kind == "h1":
            continue  # doc title line ("Methane manuscript – draft 5")
        if kind == "p" and re.match(r"^[A-Z][a-z]{2} \d+, \d{4} · @", val):
            continue  # doc byline (date chip · mention)
        if kind == "h2":
            if not title_done:
                p = heading_top(d, val); title_done = True; continue
            in_front = False
            section = val
            if val == "Declarations":
                para(d, [dict(text=val, b=False, i=False)], style="Subheadings")
            elif val == "References":
                p = d.add_paragraph(style="Subheadings")
                r = p.add_run(val); r.bold = False; r.font.size = Pt(18)
            else:
                heading_top(d, val)
            after_table = False
            continue
        if kind == "h3":
            if section == "Abstract":
                p = d.add_paragraph(style="Subheadings"); p.add_run(val)
            else:
                para(d, [dict(text=val, b=False, i=False)], style="Subheadings")
            after_table = False
            continue
        if kind == "h4":
            p = d.add_paragraph(style="Sub-subheading"); p.add_run(val)
            after_table = False
            continue
        if kind == "eq":
            segs = [dict(text=t, b=False, i=True, vert=v) for t, v in EQUATIONS[val]]
            ids = ["m_eq"] if val.startswith("y =") else (["m_h2"] if val.startswith("h^2") else (["m_t"] if val.startswith("t =") else []))
            if ids and ids[0] in ("m_h2", "m_t"):
                segs = [dict(s, hl=True) for s in segs]; found_all.add(ids[0])
            elif ids:
                segs, f = apply_highlights(segs, ids); found_all.update(f)
            para(d, segs, jc=WD_ALIGN_PARAGRAPH.CENTER, size=11.5)
            continue
        if kind == "table":
            style = "Table Grid" if last_caption in GRID_TABLES else "Plain Table 2"
            table(d, val, style, spans.HIGHLIGHT_TABLE_ROWS.get(last_caption))
            d.add_paragraph()
            after_table = True
            continue
        # paragraphs
        if val.startswith("&#91;image:") or val.startswith("[image:"):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(FIGURE1), width=Emu(5040000))
            after_table = False
            continue
        if in_front:
            text = unescape(val)
            if not d.paragraphs[-1].text.strip() or "Dermot J. Kelly" in text:
                style = "Author Names" if "Dermot J. Kelly" in text else None
            if "Dermot J. Kelly" in text:
                para(d, inline(val), style="Author Names", jc=None)
            elif re.match(r"^[¹²³]", text):
                para(d, inline(val), style="Affiliations", jc=None)
            else:
                para(d, inline(val), jc=None)
                if text.startswith("*Corresponding"):
                    d.add_paragraph()
            continue
        segs = inline(val)
        m = re.match(r"^\*\*((Table|Figure) \d+\.|S\d\.)\*\*", val)
        if section == "References":
            if CITES:
                continue  # EndNote builds the bibliography here
            mm = re.match(r"^\[(\d+)\]\s*(.*)$", unescape(val))
            p = d.add_paragraph(style="EndNote Bibliography"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run("[%s]" % mm.group(1)); r.add_tab()
            r2 = p.add_run(mm.group(2)); r2.bold = False; r2.italic = False
            continue
        if m:
            last_caption = m.group(1)
            segs, f = apply_highlights(segs, hl_ids); found_all.update(f)
            for k in f: found_count[k] = found_count.get(k, 0) + 1
            para(d, segs)
            after_table = False
            continue
        if after_table:  # table footnote
            segs, f = apply_highlights(segs, hl_ids); found_all.update(f)
            for k in f: found_count[k] = found_count.get(k, 0) + 1
            para(d, segs, size=10, base_italic=True)
            after_table = False
            continue
        segs, f = apply_highlights(segs, hl_ids); found_all.update(f)
        for k in f: found_count[k] = found_count.get(k, 0) + 1
        para(d, segs)
    d.save(out_path)
    dup = {k: v for k, v in found_count.items() if v > 1}
    if dup:
        print("WARNING highlight spans matched more than once:", dup)
    missing = set(spans.HIGHLIGHT) - found_all
    return missing


# ---------------------------------------------------------------- line numbers
def norm(s):
    s = unicodedata.normalize("NFKD", s).lower()
    return re.sub(r"[^a-z0-9]", "", s)


def to_pdf(docx_path):
    out = Path(docx_path).parent
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)],
                   check=True, capture_output=True)
    return out / (Path(docx_path).stem + ".pdf")


def numbered_lines(pdf):
    txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True).stdout
    out = []
    for ln in txt.splitlines():
        m = re.match(r"^\s{0,12}(\d{1,4})\s{2,}(\S.*)$", ln)
        if m:
            out.append((int(m.group(1)), m.group(2)))
    # keep numbers that increase in small steps (drops stray numeric table cells)
    clean, last = [], 0
    for n, t in out:
        if last < n <= last + 40:
            clean.append((n, t)); last = n
    return clean


def locate(lines, text_by_id):
    big, owner = "", []
    for n, t in lines:
        s = norm(t)
        big += s; owner += [n] * len(s)
    res = {}
    for sid, full in text_by_id.items():
        k = norm(full)
        head, tail = k[:60], k[-60:]
        a = big.find(head)
        if a < 0:
            res[sid] = None; continue
        z = big.find(tail, a)
        res[sid] = (owner[a], owner[z + len(tail) - 1]) if z >= 0 else (owner[a], owner[a])
    return res


def span_texts(md_path):
    """Full plain text of each span, taken from the manuscript markdown."""
    paras = []
    for ln in Path(md_path).read_text().splitlines():
        if ln.strip() and not ln.startswith(("|", "```", "#")):
            paras.append(plain(inline(ln)))
    eqs = {"m_eq": "+ DP + CG + a + pe + e", "m_h2": "h² = σ²ₐ / (σ²ₐ + σ²ₚₑ + σ²ₑ)",
           "m_t": "t = (σ²ₐ + σ²ₚₑ) / (σ²ₐ + σ²ₚₑ + σ²ₑ)"}
    out = {}
    for sid, (start, end) in {**spans.HIGHLIGHT, **spans.LOCATE}.items():
        if sid in eqs:
            out[sid] = eqs[sid]; continue
        for t in paras:
            a = t.find(start)
            if a >= 0:
                z = a + len(start) if end is None else t.find(end, a) + len(end)
                out[sid] = t[a:z]; break
    return out


def fmt_ranges(rngs):
    rngs = sorted(set(r for r in rngs if r))
    merged = []
    for a, b in rngs:
        if merged and a <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    return ", ".join(f"{a}" if a == b else f"{a}–{b}" for a, b in merged)


# ---------------------------------------------------------------- letter
def build_letter(md_path, out_path, refs):
    note_old = "Manuscript locations refer to draft 5 of the manuscript;"
    note_new = "Line numbers given with each response refer to the line-numbered draft 6 of the manuscript;"
    text = Path(md_path).read_text()
    lines = text.replace(note_old, note_new).splitlines()  # no-op once the doc's note is updated
    d = fresh_document()
    sect = d.element.body.find(qn("w:sectPr"))
    for ln_el in sect.findall(qn("w:lnNumType")):
        sect.remove(ln_el)
    resp_idx = 0
    used = []

    def response_suffix():
        nonlocal resp_idx
        ids = spans.LETTER_REFS[resp_idx]
        resp_idx += 1
        r = [refs.get(i) for i in ids]
        s = fmt_ranges([x for x in r if x])
        used.append((resp_idx - 1, ids, s))
        return s

    for ln in lines:
        if not ln.strip() or ln.strip() == ">":
            continue
        if ln.startswith("# "):
            heading_top(d, unescape(ln[2:])); continue
        if re.match(r"^[A-Z][a-z]{2} \d+, \d{4} · @", ln):
            continue
        if ln.startswith("## "):
            p = d.add_paragraph(); r = p.add_run(unescape(ln[3:])); r.bold = True; r.font.size = Pt(14)
            continue
        if ln.startswith("> "):
            para(d, inline(ln[2:]), base_italic=True)
            continue
        if ln.startswith("- **Line"):  # minor comment: "- **Line N:** comment **Response:** reply"
            m = re.match(r"^- \*\*(Line[^*]*)\*\*\s*(.*?)\s*\*\*Response:\*\*\s*(.*)$", ln)
            para(d, [dict(text=unescape(m.group(1)) + " ", b=True, i=True)] + [dict(s, i=True) for s in inline(m.group(2))])
            sfx = response_suffix()
            reply = m.group(3)
            if sfx and re.search(r"\((\w+), paragraph \d+\)", reply):
                # "(Background, paragraph 1)" -> "(Background, lines 46–47)"
                body = inline(re.sub(r"\((\w+), paragraph \d+\)", rf"(\1, lines {sfx})", reply))
            else:
                body = inline(reply)
                if sfx:
                    body.append(dict(text=f" (revised manuscript, lines {sfx})", b=False, i=False))
            para(d, [dict(text="Response: ", b=True, i=False)] + body)
            continue
        if ln.startswith("**Response:**"):
            sfx = response_suffix()
            segs = inline(ln)
            if sfx:
                segs.append(dict(text=f" [Revised manuscript, lines {sfx}.]", b=False, i=False))
            para(d, segs)
            continue
        if re.match(r"^\d+\. ", ln):
            para(d, inline(ln))
            continue
        para(d, inline(ln))
    assert resp_idx == len(spans.LETTER_REFS), (resp_idx, len(spans.LETTER_REFS))
    d.save(out_path)
    return used


def main():
    global CITES
    if sys.argv[1] == "--endnote":
        # Manuscript only, with EndNote temporary citations; no line numbers or letter.
        import json
        CITES = json.loads(Path(sys.argv[2]).read_text())
        ms_md, ms_out = sys.argv[3:5]
        missing = build_manuscript(ms_md, ms_out)
        if missing:
            print("WARNING highlight spans not found:", sorted(missing))
        print("written", ms_out)
        return
    ms_md, letter_md, ms_out, letter_out = sys.argv[1:5]
    missing = build_manuscript(ms_md, ms_out)
    if missing:
        print("WARNING highlight spans not found:", sorted(missing))
    pdf = to_pdf(ms_out)
    lines = numbered_lines(pdf)
    print(f"manuscript: {len(lines)} numbered lines -> {pdf}")
    refs = locate(lines, span_texts(ms_md))
    for sid, r in refs.items():
        print(f"  {sid:12s} {r}")
    used = build_letter(letter_md, letter_out, refs)
    for k, ids, s in used:
        print(f"  response {k:2d}: {s or '-'}")
    print("letter pdf:", to_pdf(letter_out))


if __name__ == "__main__":
    main()
